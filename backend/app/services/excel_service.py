import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.report_schema import ReportDataDTO, ReportType

# ─── Style helpers ─────────────────────────────────────────────────────────────
BLUE_FILL   = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
GREEN_FILL  = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
PURPLE_FILL = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid")
AMBER_FILL  = PatternFill(start_color="C65911", end_color="C65911", fill_type="solid")
PINK_FILL   = PatternFill(start_color="C0143C", end_color="C0143C", fill_type="solid")

HEADER_COLOR = {
    ReportType.summary:  BLUE_FILL,
    ReportType.activity: GREEN_FILL,
    ReportType.customer: PURPLE_FILL,
    ReportType.revenue:  AMBER_FILL,
}

REPORT_TITLE = {
    ReportType.summary:  "BÁO CÁO TỔNG HỢP",
    ReportType.activity: "BÁO CÁO HOẠT ĐỘNG KHÁCH HÀNG",
    ReportType.customer: "BÁO CÁO DANH SÁCH KHÁCH HÀNG",
    ReportType.revenue:  "BÁO CÁO DOANH THU & ĐƠN HÀNG",
}

def _auto_width(ws, headers: list[str], start_col: int = 1):
    for col_idx in range(start_col, start_col + len(headers)):
        col_letter = get_column_letter(col_idx)
        max_length = max(
            (len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[col_letter].width = max_length + 4


def _write_header(ws, data: ReportDataDTO):
    fill = HEADER_COLOR.get(data.report_type, BLUE_FILL)
    title = REPORT_TITLE.get(data.report_type, "BÁO CÁO")

    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Từ ngày: {data.start_date.strftime('%d/%m/%Y')}  —  Đến ngày: {data.end_date.strftime('%d/%m/%Y')}"
    ws["A2"].alignment = Alignment(horizontal="center")


def _write_summary_block(ws, data: ReportDataDTO, start_row: int = 4):
    s = data.summary
    avg_min = round(s.avg_duration_seconds / 60, 1)

    kpis = [
        ("Tổng lượt khách",    s.total_visitors),
        ("Khách mới",          s.new_visitors),
        ("Khách quay lại",     s.returning_visitors),
        ("Thời gian lưu trú TB", f"{avg_min} phút"),
        ("Tổng đơn hàng",     s.total_orders),
        ("Tổng doanh thu",    f"{s.total_revenue:,.0f} VND"),
    ]

    for i, (label, value) in enumerate(kpis):
        row = start_row + (i // 2)
        col = 1 if i % 2 == 0 else 4
        ws.cell(row=row, column=col, value=label).font = Font(bold=True)
        ws.cell(row=row, column=col + 1, value=value)

    return start_row + (len(kpis) // 2) + 1


def _write_daily_table(ws, data: ReportDataDTO, start_row: int):
    fill = HEADER_COLOR.get(data.report_type, BLUE_FILL)
    header_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center")

    headers = ["Ngày", "Tổng khách", "Khách mới", "Quay lại", "Thời gian TB (ph)", "Đơn hàng", "Doanh thu (VND)"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = center

    for i, row in enumerate(data.daily_stats, start_row + 1):
        ws.cell(row=i, column=1, value=row.statistic_date.strftime("%d/%m/%Y"))
        ws.cell(row=i, column=2, value=row.total_visitors)
        ws.cell(row=i, column=3, value=row.new_visitors)
        ws.cell(row=i, column=4, value=row.returning_visitors)
        ws.cell(row=i, column=5, value=round(row.avg_duration_seconds / 60, 1))
        ws.cell(row=i, column=6, value=row.total_orders)
        rev = ws.cell(row=i, column=7, value=float(row.total_revenue))
        rev.number_format = '#,##0'


def _write_customer_table(ws, data: ReportDataDTO, start_row: int):
    fill = PURPLE_FILL
    header_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center")

    headers = ["Mã khách", "Loại", "Tên khách hàng", "Lượt ghé", "Thời gian TB (ph)", "Chi tiêu (VND)", "Nhóm AI"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = center

    for i, c in enumerate(data.customers, start_row + 1):
        ws.cell(row=i, column=1, value=c.anonymous_code)
        ws.cell(row=i, column=2, value="Đã nhận diện" if c.person_type == "identified" else "Ẩn danh")
        ws.cell(row=i, column=3, value=c.customer_name or "—")
        ws.cell(row=i, column=4, value=c.total_visits)
        ws.cell(row=i, column=5, value=round(c.avg_duration_seconds / 60, 1))
        spent = ws.cell(row=i, column=6, value=float(c.total_spent))
        spent.number_format = '#,##0'
        ws.cell(row=i, column=7, value=c.segment_name or "—")


def generate_excel_report(data: ReportDataDTO) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = REPORT_TITLE.get(data.report_type, "Báo Cáo")[:31]

    _write_header(ws, data)

    # ── Nội dung theo loại ────────────────────────────────────────────────────
    if data.report_type == ReportType.summary:
        next_row = _write_summary_block(ws, data, start_row=4)
        ws.cell(row=next_row, column=1, value="CHI TIẾT THEO NGÀY").font = Font(bold=True, size=12)
        _write_daily_table(ws, data, next_row + 1)
        # Sheet 2: danh sách khách
        if data.customers:
            ws2 = wb.create_sheet("Danh sách khách hàng")
            _write_header(ws2, data)
            _write_customer_table(ws2, data, start_row=4)

    elif data.report_type == ReportType.activity:
        next_row = _write_summary_block(ws, data, start_row=4)
        ws.cell(row=next_row, column=1, value="CHI TIẾT THEO NGÀY").font = Font(bold=True, size=12)
        _write_daily_table(ws, data, next_row + 1)

    elif data.report_type == ReportType.customer:
        _write_customer_table(ws, data, start_row=4)

    elif data.report_type == ReportType.revenue:
        # Summary chỉ gồm doanh thu
        s = data.summary
        ws.cell(row=4, column=1, value="Tổng lượt khách").font = Font(bold=True)
        ws.cell(row=4, column=2, value=s.total_visitors)
        ws.cell(row=5, column=1, value="Tổng đơn hàng").font = Font(bold=True)
        ws.cell(row=5, column=2, value=s.total_orders)
        ws.cell(row=6, column=1, value="Tổng doanh thu").font = Font(bold=True)
        rev = ws.cell(row=6, column=2, value=float(s.total_revenue))
        rev.number_format = '#,##0 "VND"'
        cr = (s.total_orders / s.total_visitors * 100) if s.total_visitors else 0
        ws.cell(row=7, column=1, value="Tỷ lệ chuyển đổi").font = Font(bold=True)
        ws.cell(row=7, column=2, value=f"{cr:.1f}%")

        ws.cell(row=9, column=1, value="CHI TIẾT THEO NGÀY").font = Font(bold=True, size=12)
        _write_daily_table(ws, data, 10)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
